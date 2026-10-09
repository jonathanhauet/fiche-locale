"""
Donnees structurees (JSON-LD LocalBusiness) de toutes les agences d'un reseau, a partir de leurs fiches Google : un fichier a
remettre au developpeur du site (une ligne par agence : adresse de la page, bloc a poser, informations manquantes). Rien n'est
invente : un champ absent de la fiche Google est simplement omis et signale.
"""

import json
import re
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from urllib.parse import urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import geo_technique, google_location, google_oauth

NB_FICHES_MAX = 400
REGEX_TYPE_SCHEMA = re.compile(r"^[A-Za-z]{3,40}$")


def url_propre(url: str) -> str:
    """Adresse de la page sans parametres de suivi (« ?utm_source=GMB ») ni ancre."""
    url = (url or "").strip()
    if not url:
        return ""
    morceaux = urlsplit(url)
    return urlunsplit((morceaux.scheme, morceaux.netloc, morceaux.path, "", ""))


def collecter(db, clients: list) -> list:
    """[{"client", "infos", "categorie", "erreur"}, ...] dans l'ordre : une lecture de fiche Google par agence, en parallele."""
    resultats, travaux = [], []
    for client in clients:
        entree = {"client": client, "infos": {}, "categorie": "", "erreur": ""}
        resultats.append(entree)
        if not client.account_id or not client.location_id:
            entree["erreur"] = "Pas de fiche Google associée."
            continue
        identifiants = google_oauth.obtenir_identifiants(db, client.compte_google_id)
        if not identifiants:
            entree["erreur"] = "Compte Google non valide (à reconnecter)."
            continue
        travaux.append((entree, identifiants, client.location_id))

    def lire(travail):
        entree, identifiants, location_id = travail
        try:
            entree["infos"] = google_location.obtenir_infos_fiche(identifiants, location_id)
            entree["categorie"] = google_location.valeurs_protegees(entree["infos"]).get("categorie_nom", "")
        except Exception as erreur:
            entree["erreur"] = str(erreur)[:200]

    with ThreadPoolExecutor(max_workers=6) as pool:
        list(pool.map(lire, travaux))
    return resultats


def manques(infos: dict) -> list:
    """Informations absentes de la fiche Google, donc absentes du bloc."""
    absents = []
    if not infos.get("websiteUri"):
        absents.append("adresse de la page (site web de la fiche)")
    if not (infos.get("phoneNumbers") or {}).get("primaryPhone"):
        absents.append("téléphone")
    adresse = infos.get("storefrontAddress") or {}
    if not adresse.get("addressLines") or not adresse.get("locality"):
        absents.append("adresse postale complète")
    if not (infos.get("regularHours") or {}).get("periods"):
        absents.append("horaires")
    latlng = infos.get("latlng") or {}
    if latlng.get("latitude") is None:
        absents.append("coordonnées GPS")
    return absents


def construire_entrees(resultats: list, type_schema: str = "", reseau_nom: str = "", reseau_url: str = "") -> list:
    """Une entree par agence lue : {"agence", "url", "json_ld" (dict), "bloc" (balise script), "manques"} ; triee par nom."""
    parent = {"@type": "Organization", "name": reseau_nom.strip(), **({"url": reseau_url.strip()} if reseau_url.strip() else {})} if reseau_nom.strip() else None
    entrees = []
    for r in sorted(resultats, key=lambda r: (r["client"].nom or "").lower()):
        if not r["infos"]:
            continue
        url = url_propre(r["infos"].get("websiteUri", ""))
        donnees = geo_technique.construire_json_ld_dict(
            r["client"], r["infos"], url, r["categorie"], type_force=type_schema, parent=parent,
        )
        bloc = '<script type="application/ld+json">\n' + json.dumps(donnees, ensure_ascii=False, indent=2) + "\n</script>"
        entrees.append({"agence": r["client"].nom, "url": url, "json_ld": donnees, "bloc": bloc, "manques": manques(r["infos"])})
    return entrees


def en_json(entrees: list) -> str:
    return json.dumps(
        [{"agence": e["agence"], "url_page": e["url"], "json_ld": e["json_ld"], "informations_manquantes": e["manques"]} for e in entrees],
        ensure_ascii=False, indent=2,
    )


def en_excel(entrees: list, resultats: list) -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Blocs JSON-LD"
    feuille.append(["Agence", "Adresse de la page agence", "Bloc à coller dans la page (balise script)", "Informations absentes de la fiche Google"])
    for cellule in feuille[1]:
        cellule.font, cellule.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
        cellule.alignment = Alignment(wrap_text=True, vertical="center")
    for e in entrees:
        feuille.append([e["agence"], e["url"], e["bloc"], ", ".join(e["manques"])])
    feuille.column_dimensions["A"].width = 34
    feuille.column_dimensions["B"].width = 52
    feuille.column_dimensions["C"].width = 70
    feuille.column_dimensions["D"].width = 42
    for ligne in feuille.iter_rows(min_row=2):
        for cellule in ligne:
            cellule.alignment = Alignment(vertical="top", wrap_text=cellule.column_letter in ("C", "D"))
    feuille.freeze_panes = "B2"
    feuille.auto_filter.ref = f"A1:D{feuille.max_row}"

    non_lues = [r for r in resultats if not r["infos"]]
    if non_lues:
        notes = classeur.create_sheet("Fiches non lues")
        notes.append(["Agence", "Raison"])
        for cellule in notes[1]:
            cellule.font = Font(bold=True)
        for r in sorted(non_lues, key=lambda r: (r["client"].nom or "").lower()):
            notes.append([r["client"].nom, r["erreur"] or "Aucune donnée renvoyée par Google."])
        notes.column_dimensions["A"].width = 34
        notes.column_dimensions["B"].width = 70

    lisez = classeur.create_sheet("Mode d'emploi")
    for ligne in (
        ["Comment utiliser ce fichier"],
        ["1. Chaque ligne correspond à une page agence : collez le bloc de la colonne C dans le code de cette page (dans la balise <head> ou en bas de page)."],
        ["2. Le fichier JSON fourni en complément contient les mêmes données pour un traitement automatique (une boucle sur le modèle de page, par exemple)."],
        ["3. Vérifiez chaque bloc avec l'outil de test de Google (Rich Results Test) avant mise en ligne."],
        ["4. Une information absente (colonne D) n'est pas inventée : complétez-la dans la fiche Google, puis régénérez le fichier."],
        ["5. Les informations viennent des fiches Google des agences : si l'adresse, le téléphone ou les horaires changent, régénérez le fichier."],
    ):
        lisez.append(ligne)
    lisez["A1"].font = Font(bold=True, size=13)
    lisez.column_dimensions["A"].width = 130

    sortie = BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()
