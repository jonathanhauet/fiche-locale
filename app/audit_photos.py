"""
Audit des photos actuelles des fiches Google d'un reseau : repere les fiches dont les photos sont a corriger.
Deux niveaux : des regles objectives lues chez Google (categories, taille en pixels, anciennete) et, sur demande, une analyse
de la qualite des images par l'IA (floue, sombre, mal cadree, capture d'ecran, sans rapport avec l'activite...).
"""

import base64
import json
import os
import re
from io import BytesIO
from datetime import datetime, timezone

import requests
from anthropic import Anthropic
from openpyxl import Workbook
from PIL import Image
from openpyxl.styles import Alignment, Font, PatternFill

from . import google_business

MODELE_CLAUDE = "claude-sonnet-5"
SEUIL_PIXELS = 720                 # en dessous de ce cote minimal, la photo est trop petite pour un affichage net
NB_PHOTOS_RECOMMANDE = 10
MOIS_ANCIENNETE = 6
NB_IMAGES_ANALYSEES = 6
PAGES_MAX = 5                      # 5 pages de 100 photos : au-dela, la liste est tronquee

LIBELLES_CATEGORIE = {
    "COVER": "Couverture", "PROFILE": "Profil", "LOGO": "Logo", "EXTERIOR": "Extérieur (devanture)", "INTERIOR": "Intérieur",
    "PRODUCT": "Produit", "AT_WORK": "Au travail", "TEAMS": "Équipe", "ADDITIONAL": "Autre", "FOOD_AND_DRINK": "Nourriture", "MENU": "Menu",
}
PROBLEMES = {
    "floue": "floue", "sombre": "trop sombre", "surexposee": "surexposée", "mal_cadree": "mal cadrée", "basse_resolution": "basse résolution",
    "capture_ecran_ou_texte": "capture d'écran ou trop de texte", "photo_de_stock": "photo de banque d'images", "personnes_reconnaissables": "personnes reconnaissables",
    "sans_rapport": "sans rapport avec l'activité", "filigrane": "filigrane", "desordre": "désordre visible",
}
CONTENUS = ["devanture", "interieur", "equipe", "produit", "logo", "capture_ecran", "autre"]

SCHEMA_ANALYSE = {
    "type": "object",
    "properties": {
        "photos": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "index": {"type": "integer"},
                    "contenu": {"type": "string", "enum": CONTENUS},
                    "qualite": {"type": "integer", "description": "1 (inutilisable) a 5 (excellente)"},
                    "problemes": {"type": "array", "items": {"type": "string", "enum": list(PROBLEMES)}},
                    "commentaire": {"type": "string", "description": "Une phrase courte en francais"},
                },
                "required": ["index", "contenu", "qualite", "problemes", "commentaire"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["photos"],
    "additionalProperties": False,
}


def lire_photos(identifiants, account_id: str, location_id: str) -> tuple:
    """(toutes les photos de la fiche, True si la liste a ete tronquee a PAGES_MAX pages)."""
    photos, jeton = [], None
    for _ in range(PAGES_MAX):
        page, jeton = google_business.lister_photos_page(identifiants, account_id, location_id, jeton, 100)
        photos += page
        if not jeton:
            return photos, False
    return photos, True


def _petit_cote(photo: dict):
    largeur, hauteur = photo.get("largeur") or 0, photo.get("hauteur") or 0
    return min(largeur, hauteur) if largeur and hauteur else None


def _mois_depuis(date_iso: str):
    try:
        date = datetime.fromisoformat(date_iso.replace("Z", "+00:00"))
    except (ValueError, AttributeError):
        return None
    maintenant = datetime.now(timezone.utc)
    return (maintenant.year - date.year) * 12 + maintenant.month - date.month


def analyse_regles(photos: list, tronque: bool = False) -> dict:
    """Constats objectifs : repartition par categorie, elements manquants, taille, anciennete, et alertes en clair."""
    par_categorie = {}
    for p in photos:
        par_categorie[p["categorie"] or "ADDITIONAL"] = par_categorie.get(p["categorie"] or "ADDITIONAL", 0) + 1
    petites = [p for p in photos if (_petit_cote(p) or SEUIL_PIXELS) < SEUIL_PIXELS]
    # Google renvoie parfois la date « 1970 » (photo de profil sans date) : ignoree.
    dates = sorted((p["date_publication"] for p in photos if (p.get("date_publication") or "")[:4] > "2000"), reverse=True)
    mois = _mois_depuis(dates[0]) if dates else None
    manquants = {
        "couverture": "COVER" not in par_categorie, "logo": "LOGO" not in par_categorie and "PROFILE" not in par_categorie,
        "devanture": "EXTERIOR" not in par_categorie, "interieur": "INTERIOR" not in par_categorie,
    }
    alertes, penalite = [], 0
    if not photos:
        alertes.append("Aucune photo sur la fiche")
        penalite += 60
    else:
        if len(photos) < NB_PHOTOS_RECOMMANDE:
            alertes.append(f"Seulement {len(photos)} photo{'s' if len(photos) > 1 else ''} (minimum conseillé : {NB_PHOTOS_RECOMMANDE})")
            penalite += 15 if len(photos) >= 3 else 30
        for cle, libelle, points in (("couverture", "Pas de photo de couverture", 15), ("logo", "Pas de logo", 10),
                                     ("devanture", "Aucune photo classée Extérieur (devanture)", 15), ("interieur", "Aucune photo classée Intérieur", 5)):
            if manquants[cle]:
                alertes.append(libelle)
                penalite += points
        if petites:
            alertes.append(f"{len(petites)} photo{'s' if len(petites) > 1 else ''} trop petite{'s' if len(petites) > 1 else ''} (moins de {SEUIL_PIXELS} px)")
            penalite += 10
        if mois is not None and mois > MOIS_ANCIENNETE:
            alertes.append(f"Dernière photo ajoutée il y a {mois} mois")
            penalite += 10
    return {
        "nb": len(photos), "tronque": tronque, "par_categorie": par_categorie, "manquants": manquants, "nb_petites": len(petites),
        "derniere_photo": dates[0][:10] if dates else "", "mois_depuis_derniere": mois, "alertes": alertes, "penalite": penalite,
    }


def choisir_echantillon(photos: list, maximum: int = NB_IMAGES_ANALYSEES) -> list:
    """Photos cles a regarder : couverture, profil, logo, devanture (2 max), puis les plus recentes pour completer."""
    choisies, deja = [], set()

    def ajouter(photo):
        if photo["url"] not in deja and len(choisies) < maximum:
            deja.add(photo["url"])
            choisies.append(photo)

    for categorie, nombre in (("COVER", 1), ("PROFILE", 1), ("LOGO", 1), ("EXTERIOR", 2), ("INTERIOR", 1)):
        for photo in [p for p in photos if p["categorie"] == categorie][:nombre]:
            ajouter(photo)
    for photo in sorted(photos, key=lambda p: p.get("date_publication") or "", reverse=True):
        ajouter(photo)
    return choisies


def reduire(url: str, taille: int = 640) -> str:
    """Version reduite d'une image Google (le suffixe « =s0 » de l'adresse donne l'originale)."""
    return re.sub(r"=[swh][0-9a-z-]*$", "", url) + f"=s{taille}"


def _telecharger(url: str):
    try:
        reponse = requests.get(reduire(url), timeout=20)
    except requests.RequestException:
        return None
    if reponse.status_code != 200 or not reponse.headers.get("content-type", "").startswith("image/"):
        return None
    try:                                  # toujours en JPEG : certaines photos arrivent en PNG ou WebP
        image = Image.open(BytesIO(reponse.content)).convert("RGB")
        image.thumbnail((640, 640))
        sortie = BytesIO()
        image.save(sortie, format="JPEG", quality=85)
        return sortie.getvalue()
    except Exception:
        return None


def analyser_images(nom_agence: str, activite: str, echantillon: list) -> dict:
    """{index dans l'echantillon: {contenu, qualite, problemes, commentaire}} - vide si l'analyse est impossible."""
    cle = os.getenv("ANTHROPIC_API_KEY")
    if not cle:
        raise RuntimeError("ANTHROPIC_API_KEY manquant dans plateforme_web/.env.")
    contenu, presentes = [], []
    for indice, photo in enumerate(echantillon):
        octets = _telecharger(photo["url"])
        if not octets:
            continue
        presentes.append(indice)
        contenu.append({"type": "text", "text": f"Photo {indice} (catégorie sur la fiche Google : {LIBELLES_CATEGORIE.get(photo['categorie'], photo['categorie'] or 'Autre')})"})
        contenu.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": base64.b64encode(octets).decode()}})
    if not presentes:
        return {}
    contenu.append({"type": "text", "text": (
        f"Ces photos sont affichées sur la fiche Google de l'agence « {nom_agence} » ({activite or 'commerce de proximité'}). "
        "Pour CHAQUE photo (utilise son numéro), indique : ce qu'elle montre (devanture, intérieur, équipe, produit, logo, capture d'écran, autre), "
        "sa qualité de 1 (inutilisable) à 5 (excellente) pour une fiche professionnelle, les problèmes visibles (floue, sombre, surexposée, mal cadrée, basse "
        "résolution, capture d'écran ou trop de texte, photo de banque d'images, personnes reconnaissables, sans rapport avec l'activité, filigrane, désordre), "
        "et une phrase courte de commentaire. Sois exigeant mais juste : une photo correcte sans défaut visible mérite 4 ou 5. "
        "Une devanture doit être lisible : enseigne visible, bâtiment entier, de jour."
    )})
    reponse = Anthropic(api_key=cle).messages.create(
        model=MODELE_CLAUDE, max_tokens=2500, thinking={"type": "disabled"},
        output_config={"format": {"type": "json_schema", "schema": SCHEMA_ANALYSE}},
        messages=[{"role": "user", "content": contenu}],
    )
    texte = next((b.text for b in reponse.content if b.type == "text"), None)
    if not texte:
        return {}
    return {e["index"]: e for e in json.loads(texte)["photos"] if e["index"] in presentes}


def score_et_verdict(regles: dict, analyses: dict) -> tuple:
    """(score 0-100, a_corriger, observations IA) - la penalite des regles, plus celle des photos cles de mauvaise qualite."""
    penalite = regles["penalite"]
    observations = []
    for analyse in analyses.values():
        if analyse["qualite"] <= 2:
            penalite += 12
            observations.append(analyse["commentaire"])
        elif analyse["qualite"] == 3 and [p for p in analyse["problemes"] if p != "personnes_reconnaissables"]:
            penalite += 4          # « personnes reconnaissables » est une information (photo d'equipe), pas un defaut
    score = max(0, 100 - penalite)
    return score, score < 70, observations


def _sur(valeur, longueur: int = 500) -> str:
    """Texte sans risque pour une cellule Excel (pas de formule, longueur bornee)."""
    texte = str(valeur if valeur is not None else "").replace("\r", " ")[:longueur]
    return "\u2019" + texte if texte[:1] in ("=", "+", "-", "@") else texte


def en_excel(lignes: list) -> bytes:
    """Index Excel de l'audit : une ligne par agence (les plus a corriger d'abord), avec alertes et observations."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Audit photos"
    entetes = ["Agence", "Score photos (/100)", "À corriger", "Nombre de photos", "Couverture", "Logo ou profil", "Devanture classée", "Intérieur classé",
               "Photos trop petites", "Dernière photo", "Alertes", "Observations de l'analyse IA", "Photo de couverture", "Photo de devanture"]
    feuille.append(entetes)
    for cellule in feuille[1]:
        cellule.font, cellule.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
        cellule.alignment = Alignment(wrap_text=True, vertical="center")
    for ligne in sorted(lignes, key=lambda l: (l.get("score") if isinstance(l.get("score"), int) else 101, str(l.get("nom", "")).lower())):
        manquants = ligne.get("manquants") or {}
        cles = ligne.get("cles") or []
        feuille.append([
            _sur(ligne.get("nom")), ligne.get("score") if isinstance(ligne.get("score"), int) else "", "Oui" if ligne.get("a_corriger") else "Non",
            ligne.get("nb") if isinstance(ligne.get("nb"), int) else "",
            "Non" if manquants.get("couverture") else "Oui", "Non" if manquants.get("logo") else "Oui", "Non" if manquants.get("devanture") else "Oui",
            "Non" if manquants.get("interieur") else "Oui", ligne.get("nb_petites") or 0, _sur(ligne.get("derniere_photo")),
            _sur(" ; ".join(str(a) for a in (ligne.get("alertes") or [])), 800), _sur(" ; ".join(str(o) for o in (ligne.get("observations") or [])), 800),
            _sur(next((c.get("url") for c in cles if c.get("categorie") == "COVER"), ""), 400),
            _sur(next((c.get("url") for c in cles if c.get("categorie") == "EXTERIOR"), ""), 400),
        ])
    for lettre, largeur in zip("ABCDEFGHIJKLMN", (34, 12, 10, 12, 11, 12, 14, 12, 12, 13, 60, 60, 40, 40)):
        feuille.column_dimensions[lettre].width = largeur
    for ligne in feuille.iter_rows(min_row=2):
        for cellule in ligne:
            cellule.alignment = Alignment(vertical="top", wrap_text=cellule.column_letter in ("K", "L"))
    feuille.freeze_panes = "B2"
    feuille.auto_filter.ref = f"A1:N{feuille.max_row}"
    sortie = BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()
