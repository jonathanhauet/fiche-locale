"""
Remplacement en masse des administrateurs de fiches (voir google_admins.py) -
lit un fichier Excel prepare a partir de l'export des acces actuels
(export_acces_excel.py) : pour chaque ligne, retire l'ancien administrateur
(immediat) puis invite le nouveau (reste en attente jusqu'a acceptation par
l'adresse invitee elle-meme - aucun moyen de forcer l'acces par API).
"""

from io import BytesIO

from openpyxl import load_workbook

from . import google_admins, google_oauth, models

COLONNE_CLIENT_ID = "ID client"
COLONNE_ANCIEN_EMAIL = "Administrateur à retirer (nom affiché ou email)"
COLONNE_NOUVEL_EMAIL = "Nouvel email à inviter"
COLONNE_ROLE = "Rôle"


def lire_fichier_remplacement(octets: bytes) -> list[dict]:
    """
    Renvoie [{"client_id", "ancien_email", "nouvel_email", "role"}, ...] a
    partir d'un classeur avec les colonnes ID client / Ancien email a
    retirer / Nouvel email a inviter / Role (MANAGER par defaut si absent).
    Les lignes sans ID client ou sans nouvel email sont ignorees.
    """
    classeur = load_workbook(BytesIO(octets), read_only=True)
    feuille = classeur.active

    lignes_brutes = feuille.iter_rows(values_only=True)
    entetes = [str(valeur or "").strip() for valeur in next(lignes_brutes)]

    resultats = []
    for ligne in lignes_brutes:
        valeurs = dict(zip(entetes, ligne))
        client_id_brut = valeurs.get(COLONNE_CLIENT_ID)
        nouvel_email = str(valeurs.get(COLONNE_NOUVEL_EMAIL) or "").strip()
        if not client_id_brut or not nouvel_email:
            continue
        resultats.append({
            "client_id": int(client_id_brut),
            "ancien_email": str(valeurs.get(COLONNE_ANCIEN_EMAIL) or "").strip(),
            "nouvel_email": nouvel_email,
            "role": (str(valeurs.get(COLONNE_ROLE) or "MANAGER").strip().upper() or "MANAGER"),
        })
    return resultats


def parser_emails(texte: str) -> list[str]:
    """Emails separes par retour a la ligne, virgule ou point-virgule (voir formulaire d'invitation en masse)."""
    import re

    morceaux = re.split(r"[\n,;]+", texte or "")
    vus = set()
    emails = []
    for morceau in morceaux:
        email = morceau.strip()
        if email and email.lower() not in vus:
            vus.add(email.lower())
            emails.append(email)
    return emails


def executer_invitations_masse(db, client_ids: list[int], emails: list[str], role: str) -> list[dict]:
    """
    Invite chaque email de la liste comme administrateur sur chaque fiche
    selectionnee (produit le meme cartouche de resultat que
    executer_remplacements, pour reutiliser le meme tableau d'affichage) -
    ne retire jamais d'acces existant, contrairement a /acces/remplacer.
    """
    role = (role or "MANAGER").strip().upper() or "MANAGER"
    resultats = []
    identifiants_par_compte = {}

    clients = db.query(models.Client).filter(models.Client.id.in_(client_ids)).all()
    for client in clients:
        for email in emails:
            resultat = {
                "client_id": client.id, "client_nom": client.nom, "ancien_email": "", "statut_retrait": "",
                "nouvel_email": email,
            }

            if not client.location_id:
                resultat["statut_invitation"] = "Fiche Google non associée"
                resultats.append(resultat)
                continue

            if client.compte_google_id not in identifiants_par_compte:
                identifiants_par_compte[client.compte_google_id] = google_oauth.obtenir_identifiants(db, client.compte_google_id)
            identifiants = identifiants_par_compte[client.compte_google_id]

            if not identifiants:
                resultat["statut_invitation"] = "Compte Google non valide pour ce client"
                resultats.append(resultat)
                continue

            try:
                google_admins.inviter_administrateur(identifiants, client.location_id, email, role)
                resultat["statut_invitation"] = "Invitation envoyée (en attente d'acceptation par cette adresse)"
            except Exception as erreur:
                resultat["statut_invitation"] = f"Erreur : {erreur}"

            resultats.append(resultat)

    return resultats


def executer_remplacements(db, lignes: list[dict]) -> list[dict]:
    """
    Execute chaque remplacement et renvoie la liste enrichie avec
    "client_nom", "statut_retrait", "statut_invitation" - une ligne en erreur
    n'empeche jamais les autres de s'executer.
    """
    resultats = []
    identifiants_par_compte = {}

    for ligne in lignes:
        client = db.get(models.Client, ligne["client_id"])
        resultat = {**ligne, "client_nom": client.nom if client else "?", "statut_retrait": "", "statut_invitation": ""}

        if not client or not client.location_id:
            resultat["statut_invitation"] = "Client introuvable ou sans fiche Google associée"
            resultats.append(resultat)
            continue

        if client.compte_google_id not in identifiants_par_compte:
            identifiants_par_compte[client.compte_google_id] = google_oauth.obtenir_identifiants(db, client.compte_google_id)
        identifiants = identifiants_par_compte[client.compte_google_id]

        if not identifiants:
            resultat["statut_invitation"] = "Compte Google non valide pour ce client"
            resultats.append(resultat)
            continue

        if ligne["ancien_email"]:
            try:
                admins_actuels = google_admins.lister_administrateurs(identifiants, client.location_id)
                # "ancien_email" ne correspondra que si c'est encore une
                # invitation en attente (email tel quel) ou si l'export a ete
                # prepare avec le NOM AFFICHE (voir google_admins.py) - Google
                # n'expose plus l'email une fois l'invitation acceptee.
                cible = next(
                    (a for a in admins_actuels if a["identifiant"].strip().lower() == ligne["ancien_email"].strip().lower()),
                    None,
                )
                if cible:
                    google_admins.retirer_administrateur(identifiants, cible["nom_ressource"])
                    resultat["statut_retrait"] = "Retiré"
                else:
                    resultat["statut_retrait"] = "Introuvable sur cette fiche (déjà retiré ?)"
            except Exception as erreur:
                resultat["statut_retrait"] = f"Erreur : {erreur}"

        try:
            google_admins.inviter_administrateur(identifiants, client.location_id, ligne["nouvel_email"], ligne["role"])
            resultat["statut_invitation"] = "Invitation envoyée (en attente d'acceptation par cette adresse)"
        except Exception as erreur:
            resultat["statut_invitation"] = f"Erreur : {erreur}"

        resultats.append(resultat)

    return resultats
