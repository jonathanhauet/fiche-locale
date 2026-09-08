"""
Export Excel des administrateurs actuels des fiches (qui a acces a quoi) -
voir google_admins.py et /acces dans main.py. Une ligne par administrateur
(une fiche peut en avoir plusieurs) ; sert aussi de base pour preparer le
fichier de remplacement en masse (colonne "ID client" reutilisable telle
quelle dans le fichier importe par /acces/remplacer).
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import google_admins, google_oauth

ENTETES = [
    "ID client", "Nom", "Étiquettes",
    "Administrateur (nom affiché, ou email si invitation en attente)",
    "ID compte (référence, pas un email)", "Rôle", "Invitation en attente", "Erreur",
]


def generer_export(db, clients: list) -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Accès actuels"
    feuille.append(ENTETES)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)

    identifiants_par_compte = {}

    for client in clients:
        etiquettes = ", ".join(e.nom for e in client.etiquettes)

        if not client.account_id or not client.location_id:
            feuille.append([client.id, client.nom, etiquettes, "", "", "", "", "Pas de fiche Google associée"])
            continue

        if client.compte_google_id not in identifiants_par_compte:
            identifiants_par_compte[client.compte_google_id] = google_oauth.obtenir_identifiants(db, client.compte_google_id)
        identifiants = identifiants_par_compte[client.compte_google_id]

        if not identifiants:
            feuille.append([client.id, client.nom, etiquettes, "", "", "", "", "Compte Google non valide"])
            continue

        try:
            admins = google_admins.lister_administrateurs(identifiants, client.location_id)
        except Exception as erreur:
            feuille.append([client.id, client.nom, etiquettes, "", "", "", "", str(erreur)])
            continue

        if not admins:
            feuille.append([client.id, client.nom, etiquettes, "(aucun administrateur trouvé)", "", "", "", ""])
            continue

        for admin in admins:
            feuille.append([
                client.id, client.nom, etiquettes,
                admin["identifiant"],
                admin["compte_id"],
                google_admins.LIBELLES_ROLE.get(admin["role"], admin["role"]),
                "Oui" if admin["en_attente"] else "",
                "",
            ])

    for indice, entete in enumerate(ENTETES, start=1):
        feuille.column_dimensions[get_column_letter(indice)].width = max(16, len(entete) + 2)
    feuille.freeze_panes = "A2"

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
